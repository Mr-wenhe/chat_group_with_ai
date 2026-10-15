from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = Path("docs/autonomous_group_work/p3-evidence/generated")


def set_cell_text(cell, value, *, white=False):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_font(run)
            run.font.size = Pt(8)
            if white:
                run.font.color.rgb = RGBColor(255, 255, 255)


def set_font(run):
    run.font.name = "PingFang SC"
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    fonts.set(qn("w:ascii"), "PingFang SC")
    fonts.set(qn("w:hAnsi"), "PingFang SC")
    fonts.set(qn("w:eastAsia"), "PingFang SC")
    fonts.set(qn("w:hint"), "eastAsia")
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    language.set(qn("w:eastAsia"), "zh-CN")


def set_document_fonts(document):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_font(run)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_font(run)


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def build(path, *, corrected):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = document.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "PingFang SC")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "PingFang SC")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")

    title = document.add_heading("桌面清单 App 上线审查报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("版本：2.4.0　审查日期：2026 年 10 月 12 日　负责人：林青")
    document.add_heading("审查摘要", level=1)
    document.add_paragraph(
        "本次检查窗口自适应、跨页列表、搜索筛选、键盘操作和重启后的状态保存。"
        "审查对象为内部候选版，问题按复现条件记录，修复后逐项复测。"
    )
    document.add_heading("关键观察", level=1)
    spelling = "界面缩放后，筛选条件应保持不变。" if corrected else "界面缩入后，筛选条件应保持不变。"
    document.add_paragraph(spelling)
    if corrected:
        document.add_paragraph("本次审查范围：窗口缩放、跨页列表、筛选与状态保存。")
    else:
        document.add_paragraph("请在此处填写项目摘要。")
    document.add_paragraph(
        "初步结果显示，名单分页边界需要重点核对；快速连续切换筛选条件时，"
        "列表总数与可见行数必须保持一致。"
    )

    document.add_page_break()
    document.add_heading("验收记录（跨页表格）", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headings = ["编号", "操作场景", "预期结果", "实际记录"]
    for cell, heading in zip(table.rows[0].cells, headings):
        set_cell_text(cell, heading)
    for index in range(1, 65):
        if index == 45:
            continuation = table.add_row()
            for cell, heading in zip(continuation.cells, headings):
                set_cell_text(cell, heading)
            repeat_header(continuation)
            continuation.cells[0].paragraphs[0].paragraph_format.page_break_before = True
        row = table.add_row()
        row_data = [
            f"AC-{index:02d}",
            f"打开名单并切换筛选 {index}",
            "总数、可见行和当前页一致",
            "保存后重开，条件仍保留",
        ]
        for column, (cell, value) in enumerate(zip(row.cells, row_data)):
            hide_cell = not corrected and index == 18 and column == 3
            set_cell_text(cell, value, white=hide_cell)

    document.add_page_break()
    document.add_heading("问题与复测结论", level=1)
    document.add_paragraph(
        "复现：在窗口缩放后快速切换筛选，再关闭并重开。预期筛选项和页码恢复一致；"
        "实际结果应逐项记录，不能仅依据成功提示。"
    )
    document.add_paragraph(
        "复测记录：修复后重新执行相同动作，并核对持久化文件中的筛选值及页码。"
    )
    conclusion = (
        "结论：排版检查已完成，待独立角色复测后签署。"
        if corrected
        else "结论：待完成全部页面排版检查和独立角色复测后签署。"
    )
    document.add_paragraph(conclusion)
    set_document_fonts(document)
    document.save(path)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    build(OUTPUT / "review-before.docx", corrected=False)
    build(OUTPUT / "review-after.docx", corrected=True)
    print(f"created={OUTPUT.resolve()}")


if __name__ == "__main__":
    main()
